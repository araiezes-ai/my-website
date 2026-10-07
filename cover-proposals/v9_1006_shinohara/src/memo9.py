from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.oxml.ns import qn
d=Document(); s=d.sections[0]; s.page_width=Mm(210); s.page_height=Mm(297)
for m in ('left_margin','right_margin'): setattr(s,m,Mm(16))
s.top_margin=Mm(14); s.bottom_margin=Mm(14)
st=d.styles['Normal']; st.font.name='Yu Gothic'; st.font.size=Pt(8.5); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Yu Gothic')
def H(t,l=1):
    p=d.add_heading(t,l)
    for r in p.runs: r.font.name='Yu Gothic'; r.element.rPr.rFonts.set(qn('w:eastAsia'),'Yu Gothic'); r.font.color.rgb=RGBColor(0x23,0x27,0x2c)
def P(t): p=d.add_paragraph(t); p.paragraph_format.space_after=Pt(3)
H('台割 案1（篠原修正案）　手書き修正の反映表',0)
P('10/6の修正指示（手書き）をすべて反映し、修正案のデザイン（すりガラスの球、波形・円形の写真の切り抜き、ノドを越える金の軌道線、大きな章番号の白抜き数字）を再現しました。凡例：Ⓐ＝タイトル下の和文サブコピー、Ⓑ＝本文、Ⓒ＝見出し横のグレーの説明文。')
rows=[['ページ','手書きの指示','反映内容'],
['共通','セクション見出し（OUR JOURNEY等）のフォントサイズをⒶと同じに（以下同）','OUR JOURNEY／OUR RESOURCES／TRADING HUB／OWNERSHIP／JOINT VENTURES等の英字見出しと横の和文をⒶと同じ8.2ptに統一'],
['P.04','「資源投資」「トレーディング」を目立たせる','両語を太字・色付き（投資＝金、トレーディング＝青）で大きく表示'],
['P.04','2つの輪が軸であることを強調／中央に「Value chain」','金（投資）と青（トレーディング）の2本の輪を太く描き、中央に「Value chain／バリューチェーン」を配置。5つの段階は輪の上に配置'],
['P.05','Ⓑと同じ本文を入れる','タイトル下に本文を追加'],
['P.05','OUR RESOURCESに「鉄鉱石から電池資源・肥料・アルミまで世界に広がる事業領域」を追加（Ⓒと同じ色）','見出し横に追加。旧「主力の2事業は次の見開きで」は削除'],
['P.05','肥料資源（Fertilizer resource）を追加。写真がないので丸だけ','写真なしの金の輪で追加'],
['P.05','TRADING HUB「世界の資源と需要をつなぐ販売網」：RtM International／Japan／Bharat／Europe／Americas（丸は同サイズ）＋中国・UAE・インドネシア・タイ・チリ（やや小さく）','5社を同サイズの円、5拠点を小さめの円で一列に配置（P.08の拠点一覧をここへ移動）'],
['P.06','BMAはMDPとBHPの50:50出資の合弁会社と分かる図を入れる（三菱商事→100%→MDP 50%／BHP 50%→BMA（BHP Mitsubishi Alliance））','出資構成図を追加'],
['P.07','本文（Ⓑ）を入れる：パートナー企業が多い、規模が大きく、幅広い関係を構築している、声がかかる存在である','右段の主役を本文（リード＋本文）にし、強み01〜04は下段へ移動'],
['P.08','Ⓑの文を入れる','数字の下に本文を追加'],
['P.08','グローバルネットワーク10拠点（丸で囲み）','P.05のTRADING HUBへ移動したため削除'],
['P.08','トレーディング事業例 Copper Value Chain：鉱山会社→トレーディング(RtM)→製鉄・製錬→トレーディング(RtM)→主な最終用途（例）Base Demand／Energy Transition／AI & Data Centers','下段に流れ図として追加'],
['P.09','サブコピーの「銅・クリティカルミネラルを中心に、原料炭まで」を削除→「加工技術、回収、リサイクル…」','「加工技術・回収・リサイクルなど、業界のボトルネックに挑む技術へ投資する」に変更'],
['P.09','CiDRA：スタートアップ／浮遊（選鉱）の高度化、Jetti：スタートアップ／硫化鉱のリーチング化、DESCycle：銅スクラップの回収強化、COAL〇〇〇〇は削除','3例に整理し、それぞれ見出しに追加。原料炭の枠は削除'],
['P.10','本文を入れる','サブコピーの下に本文を追加'],
['P.10','JV実績の見出しサイズⒶと同じ／社名を1行に','見出しを拡大し、BHP・Rio Tinto・Anglo Americanを1行に'],
['P.10','パートナーに評価される理由：中長期的な視座／リスクシェアリング／総合力','3つの項目を最下段に追加']]
t=d.add_table(rows=len(rows),cols=3); t.style='Light Grid Accent 1'
for i,r in enumerate(rows):
    for j,c in enumerate(r):
        cell=t.cell(i,j); cell.text=c; cell.width=Mm([16,82,80][j])
        for p in cell.paragraphs:
            for rr in p.runs: rr.font.size=Pt(7.8); rr.bold=(i==0)
d.add_paragraph()
P('手書きの指示がなかったP.02–03・P.11は修正案のデザインのまま組み直しました（P.11は外部環境の3つのすりガラス球と、最後に現れる日本側の地球儀）。本文は英文のダミーです。地図上の拠点位置は仮置きです。')
d.save('v9/台割案1_篠原修正案_反映表.docx')
