from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
FS=8.6
doc=Document(); s=doc.sections[0]
s.page_width=Mm(210); s.page_height=Mm(297); s.left_margin=s.right_margin=Mm(14); s.top_margin=Mm(11); s.bottom_margin=Mm(9)
st=doc.styles['Normal']; st.font.name='Yu Gothic'; st.font.size=Pt(FS); st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Yu Gothic')
st.paragraph_format.space_after=Pt(0); st.paragraph_format.line_spacing=1.12
def P(t='',size=None,bold=False,color=None,after=0,before=0,par=None):
    p=par or doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before)
    if t:
        r=p.add_run(t); r.bold=bold
        if size: r.font.size=Pt(size)
        if color: r.font.color.rgb=RGBColor.from_string(color)
    return p
def LBL(cell,label,text,after=1.5):
    q=cell.add_paragraph(); q.paragraph_format.space_after=Pt(after)
    r=q.add_run(label+'　'); r.bold=True; r.font.color.rgb=RGBColor.from_string('1F3A40'); q.add_run(text)
def BUL(cell,items,after=1.5):
    for t in items:
        q=cell.add_paragraph(); q.paragraph_format.left_indent=Mm(3); q.paragraph_format.first_line_indent=Mm(-3); q.paragraph_format.space_after=Pt(0.5)
        q.add_run('・'+t)
    cell.paragraphs[-1].paragraph_format.space_after=Pt(after)
def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexc); tcPr.append(sh)
P('金属資源グループ 会社案内（英語版）',7.6,color='666666')
P('表紙デザイン3案　制作の意図と見どころ',13.5,True,after=2)
P('3案共通の軸：海外の商談相手（ニューヨーク・ロンドンのビジネスパーソンを想定）が手に取り、「ひと目で金属資源の会社と分かる」「白を基調とした品格」「日本企業らしさ」の3点を満たすこと。3案とも、白地・細身の書体・色数を絞った構成とし、違いは「何で金属資源を語るか（素材／地球／現場）」に置いています。',FS,after=4)
cases=[
('案①　鋼のヘアライン（Steel Hairline）',
 '鋼板の表面仕上げ「ヘアライン（細い筋目）」で世界地図を描き、実際の事業拠点・資産22か所から細い線が上へ立ち上がる構成。コアメッセージ「Strength to Hold. Power to Connect.」を、写真に頼らず“素材”と“つながり”で表現しました。',
 ['素材を「描く」のではなく「質感として見せる」：線は極細にし、左右方向に明暗を変えることで、光を受けた鋼板の反射を再現。印刷では銀インキや箔押しに置き換えられる設計です。',
  '縦長の紙面を「上昇」に使う：地図を紙面の下に置いて“大地”とし、線を上へ伸ばすことで、資源が世界から立ち上がり、つながっていく動きを生みました。',
  '視線の流れを設計：タイトルの真下では線を止めて可読性を守り、右側だけ線を高く伸ばすことで、視線が「タイトル → 朱の線 → 日本 → 世界」へと自然に降りていきます。',
  '色は「情報」としてだけ使う：朱は日本と東京から伸びる1本のみ。装飾ではなく「起点はここ」という意味を持たせた一点です。'],
 '図案そのものが「鋼（Strength）」と「つながり（Connect）」を同時に語る／線の起点は中面と同じ実在の拠点で、誇張のない“事実でできた絵”／日の丸を思わせる控えめな和'),
('案②　ガラスの地球儀（Glass Globe）',
 'アジア・オセアニアを中心にした透明感のある地球儀と、その周りを巡る金色の軌道線で、「日本を起点に世界の資源と需要をつなぐ」姿を表現。大陸（国土）は、鉱物の輝きを想起させるシャンパンゴールドと、鉱石の深みを思わせるグラファイトグレーで塗り分け、地球儀そのものが「金属資源」を語るようにしました。',
 ['地球の「向き」を選ぶ：一般的な大西洋中心ではなく、日本と豪州・南米の資源国の関係が見えるアジア・オセアニア中心の角度を選び、事業の地理を一枚で伝えます。',
  '素材で姿勢を語る：ガラスの透明感は、コアメッセージにある「フェアな立場」「多様なパートナーとの間に立つ」姿勢の暗示でもあります。',
  '色の質感にこだわる：金は派手さを抑えたシャンパンゴールド、グレーは鉱物（黒鉛）の名を持つグラファイト。高級感を出しつつ、誇示しないトーンに収めています。',
  '軌道線＝循環と持続：地球を巡る細い楕円は、資源の流れと「持続的な安定供給」というグループミッションを、言葉を使わずに表します。'],
 '最も分かりやすく、誰に渡しても誤解のない明快さ／白地と広い余白で明るく開かれた印象／金色の軌道線を中面の地図や図解に展開でき、冊子全体を統一しやすい'),
('案③　鉱山のモダンアート（Mine as Modern Art）',
 '貴社サイトに掲載の露天掘り鉱山の写真を縦の短冊に切り、上下にずらして再構成。実際の事業の写真を使いながら、1枚の現代美術作品のように見せ、商談の場で最も強い第一印象を残す案です。',
 ['「記録写真」を「作品」に変える：切ってずらすという一手間だけで、見慣れた鉱山写真が抽象画のように見え、思わず手が止まる表紙になります。',
  '鉱山の構造と呼応させる：短冊の段差は、露天掘りの段々（ベンチ）や地層の重なりと響き合い、長い時間をかけて資源を育てる事業の時間軸も感じさせます。',
  '規則と不規則のバランス：短冊は等幅のグリッドで揃え、ずれ幅だけを不規則にすることで、整然さとリズムを両立させています。',
  '朱の一点で締める：段差の足元に小さな朱の正方形を置き、画面の重心と「起点」を示しています。中面でも同じ手法を繰り返し、シリーズとしての一体感を出せます。'],
 '写真でありながら抽象画のような強さ／露天掘りの圧倒的なスケールが、そのまま事業規模を伝える／中面の写真まで一貫した表現に展開できる'),
]
tbl=doc.add_table(rows=len(cases),cols=1)
for i,(t,why,dv,best) in enumerate(cases):
    c=tbl.rows[i].cells[0]; shade(c,'F3F4F5' if i%2==0 else 'FFFFFF')
    p=c.paragraphs[0]; r=p.add_run(t); r.bold=True; r.font.size=Pt(9.8); p.paragraph_format.space_after=Pt(1.5)
    LBL(c,'制作の意図',why)
    q=c.add_paragraph(); r=q.add_run('デザイナーの視点'); r.bold=True; r.font.color.rgb=RGBColor.from_string('1F3A40'); q.paragraph_format.space_after=Pt(0.5)
    BUL(c,dv)
    LBL(c,'見どころ',best,after=3)
P(before=3)
P('「あなたならどれを選びますか？」と聞かれたら',9.8,True,'1F3A40',after=1)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); r=p.add_run('「私は案①です。3案の中で唯一、表紙だけでコアメッセージを言い切れているからです。」'); r.bold=True
for h,b in [('① メッセージと一致している。','“Strength to Hold”を鋼の質感で、“Power to Connect”を拠点から立ち上がる線で表し、冊子を開く前から主張が伝わります。'),
            ('② 事実でできていて、誇示がない。','線の起点は中面と同じ実在の22拠点。ストーリーラインの「優位性を誇示しない」方針に、表紙の段階から沿っています。'),
            ('③ 投資と販売の“両輪”に偏りがない。','鉱山写真は川上の印象が強くなりますが、地図と線の表現は投資と販売を含む事業全体を等しく表せます。'),
            ('④ 長く使え、本制作で品質が上がる。','写真の許諾や解像度に左右されず、銀・箔・特色の朱といった印刷加工で完成度を高められます。')]:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0.5); r=p.add_run(h); r.bold=True; p.add_run(b)
P('ひと言で添えるなら：「案②は“分かりやすさ”、案③は“インパクト”。案①は、貴社の“らしさ”です。」',FS,True,'1F3A40',before=2,after=3)
P('制作プロセスについて：本3案は、表現の方向性を短期間で幅広く検討するためにAIを活用して制作しました。ご選定いただいた案は、デザイナーが書体・色・印刷仕様まで設計し直し、本制作として仕上げます。',7.8,color='555555')
doc.save('cover_explanation.docx')
